import sys
from docx import Document
from docx.oxml.shared import OxmlElement
from docx.oxml.ns import qn

def delete_paragraph(paragraph):
    p = paragraph._element
    p.getparent().remove(p)
    paragraph._p = paragraph._element = None

def insert_toc(paragraph):
    run = paragraph.add_run()
    
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = ' TOC \\\\o "1-3" \\\\h \\\\z \\\\u '
    
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

base_path = r'\\wsl.localhost\Ubuntu\home\tonilogar\trabajos\programming\tonilogar.github.io\tfb\delivery\000_entrega_03_antonio_lopez_018_numerado.docx'
out_path = r'\\wsl.localhost\Ubuntu\home\tonilogar\trabajos\programming\tonilogar.github.io\tfb\delivery\000_entrega_03_antonio_lopez_018_final.docx'

doc = Document(base_path)

index_start = -1
index_end = -1
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if 'ndice interactivo' in text or 'Índice interactivo' in text:
        index_start = i
    if index_start != -1 and i > index_start:
        if '11. Referencias Bibliogr' in text:
            index_end = i
            break
        if text and not (text.startswith('Glosario') or text[0].isdigit()):
            index_end = i - 1
            break

if index_start != -1 and index_end != -1:
    p_index = doc.paragraphs[index_start]
    p_index.clear()
    
    r = p_index.add_run('Índice interactivo')
    r.bold = True
    
    p_toc = doc.paragraphs[index_start + 1]
    p_toc.clear()
    insert_toc(p_toc)
    
    for i in range(index_start + 2, index_end + 1):
        delete_paragraph(doc.paragraphs[i])
        
    settings = doc.settings.element
    update_fields = OxmlElement('w:updateFields')
    update_fields.set(qn('w:val'), 'true')
    settings.append(update_fields)
    
    doc.save(out_path)
    print('Native TOC inserted successfully.')
else:
    print('Could not find index boundaries.')
