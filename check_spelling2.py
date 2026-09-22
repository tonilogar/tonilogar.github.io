from spellchecker import SpellChecker
from docx import Document
import re
from collections import Counter

doc_path = r'\\wsl.localhost\Ubuntu\home\tonilogar\trabajos\programming\tonilogar.github.io\tfb\delivery\entrega_deposito_final_primera_defensa_antonio_lopez.docx'

doc = Document(doc_path)
spell = SpellChecker(language='es')

text = ' '.join(p.text for p in doc.paragraphs)
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            text += ' ' + cell.text

words = re.findall(r'\b[a-zA-Z·ÈÌÛ˙¡…Õ”⁄Ò—]+\b', text)

unknown = spell.unknown(words)

counts = Counter([w for w in words if w in unknown])

print("Top misspelled words:")
for word, count in counts.most_common(100):
    print(f"{word}: {count}")
